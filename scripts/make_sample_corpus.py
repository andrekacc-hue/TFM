"""Genera el corpus de ejemplo (ficticio) en samples/corpus: un PDF, un DOCX y un Markdown.

Requiere reportlab y python-docx (dependencias dev del ingest-service).

Uso:
    python scripts/make_sample_corpus.py
"""

from __future__ import annotations

import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "samples" / "corpus"

HSM_PAGES = [
    (
        "PROC-SEG-021 Renovación de certificados del HSM",
        [
            "1. Alcance",
            "Este procedimiento describe la renovación anual del certificado TLS que protege la",
            "comunicación entre el switch de pagos y el módulo de seguridad hardware (HSM).",
            "",
            "2. Prerrequisitos",
            "- Ventana de mantenimiento aprobada por el Comité de Cambios (mínimo 72 horas antes).",
            "- Presencia de dos custodios de claves (control dual). Ninguna persona sola puede",
            "  completar el procedimiento.",
            "- Solicitud de firma (CSR) generada dentro del HSM: la clave privada nunca sale del equipo.",
        ],
    ),
    (
        None,
        [
            "3. Pasos de renovación",
            "Paso 1: El custodio A inicia sesión en la consola del HSM y genera el CSR con la plantilla",
            "TLS-SWITCH-2026. El custodio B verifica el nombre común y la longitud de clave (RSA 3072).",
            "Paso 2: Enviar el CSR a la Autoridad Certificadora interna mediante el portal de PKI y",
            "registrar el número de solicitud en el ticket de cambio.",
            "Paso 3: Importar el certificado firmado en el HSM y en el almacén de confianza del switch.",
            "Paso 4: Reiniciar únicamente el servicio de canal seguro del switch, nunca el HSM completo.",
        ],
    ),
    (
        None,
        [
            "4. Verificación y reversión",
            "Verificar con una transacción de prueba de bajo monto que el canal negocia TLS 1.2 o superior",
            "con el nuevo certificado. Si la prueba falla, restaurar el certificado anterior, que debe",
            "permanecer en el HSM hasta su fecha de expiración, y escalar a Seguridad de la Información.",
            "Registrar la evidencia de ambos custodios en el ticket de cambio.",
        ],
    ),
]

DB_LOCKS_SECTIONS = [
    ("Objetivo", "Guiar al operador de TI Producción en la detección y gestión de bloqueos en la base de datos "
                 "del core bancario durante el horario de atención y los procesos nocturnos."),
    ("Detección", "Un bloqueo se identifica cuando el panel de monitoreo muestra sesiones en espera por más de "
                  "120 segundos sobre las tablas de cuentas o movimientos. La alerta BD-LOCK-01 se dispara "
                  "automáticamente y abre un evento en la consola de operaciones."),
    ("Acciones permitidas al operador", "El operador puede identificar la sesión bloqueante, el usuario de "
                                        "aplicación y el programa que la originó, y documentarlos en el evento. "
                                        "El operador NO debe terminar sesiones de base de datos."),
    ("Escalamiento al DBA", "Si el bloqueo supera los 5 minutos o afecta a transacciones en línea, escalar al "
                            "DBA de turno por la línea de guardia. Solo el DBA, con aprobación del responsable "
                            "de la aplicación, puede terminar la sesión bloqueante."),
]

VPN_FAQ = """# FAQ-VPN-001 Acceso remoto para guardias de TI

## ¿Quién puede usar la VPN de guardia?

Solo el personal de TI Producción incluido en el rol de guardia vigente. El acceso se habilita
automáticamente al inicio del turno y se revoca al finalizarlo.

## ¿Qué hago si la VPN rechaza mi conexión con el error VPN-403?

El error **VPN-403** indica que tu usuario no está en el rol de guardia activo. Verifica en el
calendario de guardias que el turno esté asignado a tu nombre. Si es correcto, solicita al
coordinador de guardia que actualice el rol; el cambio se aplica en un máximo de 10 minutos.

## ¿Y si el error es VPN-TIMEOUT?

Indica que el concentrador VPN no respondió. Comprueba tu conexión a internet y reintenta con
el perfil secundario `vpn-guardia-2`. Si ambos perfiles fallan, reporta el incidente a la Mesa
de Servicio indicando la hora y el perfil utilizado.

## ¿Puedo compartir mi segundo factor (MFA) con otro guardia?

No. El segundo factor es personal e intransferible y su uso compartido es una falta grave según
la política de seguridad de la información.
"""

MANIFEST = [
    {"file": "PROC-SEG-021_renovacion_certificados_HSM.pdf", "title": "PROC-SEG-021 Renovación de certificados del HSM",
     "source_type": "procedimiento", "area": "operaciones"},
    {"file": "MAN-BD-003_gestion_bloqueos_bd.docx", "title": "MAN-BD-003 Gestión de bloqueos en la base de datos del core",
     "source_type": "manual", "area": "operaciones"},
    {"file": "FAQ-VPN-001_acceso_remoto_guardias.md", "title": "FAQ-VPN-001 Acceso remoto VPN para guardias de TI",
     "source_type": "faq", "area": "general"},
]


def make_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle(HSM_PAGES[0][0] or "")
    for title, lines in HSM_PAGES:
        y = 790
        if title:
            pdf.setFont("Helvetica-Bold", 14)
            pdf.drawString(50, y, title)
            y -= 30
        pdf.setFont("Helvetica", 10)
        for line in lines:
            pdf.drawString(50, y, line)
            y -= 15
        pdf.showPage()
    pdf.save()


def make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("MAN-BD-003 Gestión de bloqueos en la base de datos del core", level=0)
    for heading, text in DB_LOCKS_SECTIONS:
        document.add_heading(heading, level=1)
        document.add_paragraph(text)
    document.add_heading("Matriz de contacto", level=1)
    table = document.add_table(rows=3, cols=3)
    for r, row in enumerate([("Rol", "Canal", "Tiempo de respuesta"),
                             ("DBA de turno", "Línea de guardia", "15 minutos"),
                             ("Responsable de aplicación", "Correo y llamada", "30 minutos")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    make_pdf(OUT / MANIFEST[0]["file"])
    make_docx(OUT / MANIFEST[1]["file"])
    (OUT / MANIFEST[2]["file"]).write_text(VPN_FAQ, encoding="utf-8")
    (OUT / "manifest.json").write_text(json.dumps(MANIFEST, ensure_ascii=False, indent=2), encoding="utf-8")
    for item in MANIFEST:
        print(f"  {item['file']}  ({(OUT / item['file']).stat().st_size} bytes)")
