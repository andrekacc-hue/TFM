from __future__ import annotations

import io
from typing import Any
from uuid import uuid4

import fakeredis
import pytest

from app.config import IngestSettings
from app.core.chunker import RecursiveChunker
from app.core.embedder import BatchEmbedder
from app.core.pipeline import IngestionPipeline
from app.core.queue import IngestQueue
from shared.exceptions import NotFoundError

INTERNAL_KEY = "clave-interna-de-pruebas-123"

PROCEDURE_MD = """# PROC-OPS-014 Reinicio del batch

## Paso 1 — Diagnóstico

Consultar las sesiones **bloqueantes** sobre las tablas de saldos. Ver [guía](http://intranet/guia).

## Paso 2 — Relanzamiento

```
restart --from-checkpoint
```

| Código | Significado |
|--------|-------------|
| CIE-201 | Timeout |
"""


@pytest.fixture
def settings() -> IngestSettings:
    return IngestSettings(
        _env_file=None,
        supabase_url="https://demo.supabase.co",
        supabase_service_role_key="sb_secret_test",
        supabase_anon_key="sb_publishable_test",
        internal_api_key=INTERNAL_KEY,
        redis_url="redis://fake",
        preload_models=False,
        worker_enabled=False,
        log_format="console",
        worker_block_ms=100,
    )


# ----------------------------------------------------------------------
# Generadores de documentos reales
# ----------------------------------------------------------------------
def make_pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            pdf.drawString(50, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Introducción sin encabezado: alcance del manual de contingencia ATM.")
    document.add_heading("Síntomas", level=1)
    document.add_paragraph("Los retiros son rechazados con el código 91 en varios cajeros a la vez.")
    document.add_heading("Escalamiento", level=2)
    document.add_paragraph("Si no se restablece en 15 minutos, escalar al proveedor del switch.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Nivel", "Contacto"
    table.cell(1, 0).text, table.cell(1, 1).text = "N2", "Proveedor switch"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------
# Dobles de prueba
# ----------------------------------------------------------------------
class FakeEmbedder:
    model_name = "fake-embedder"
    dimension = 4

    def __init__(self) -> None:
        self.batches: list[list[str]] = []

    async def embed_query(self, text: str) -> list[float]:
        return [0.1, 0.2, 0.3, 0.4]

    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.batches.append(list(texts))
        return [[float(len(t) % 7), 0.0, 0.0, 1.0] for t in texts]

    def load(self) -> None:
        return None


class FakeStorage:
    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}
        self.fail_download: Exception | None = None

    async def upload(self, bucket: str, path: str, content: bytes, *, content_type: str = "", upsert: bool = False) -> str:
        self.objects[f"{bucket}/{path}"] = content
        return path

    async def download(self, bucket: str, path: str) -> bytes:
        if self.fail_download:
            raise self.fail_download
        key = f"{bucket}/{path}"
        if key not in self.objects:
            raise NotFoundError("Objeto no encontrado en Storage")
        return self.objects[key]

    async def remove(self, bucket: str, paths: list[str]) -> None:
        for path in paths:
            self.objects.pop(f"{bucket}/{path}", None)

    async def aclose(self) -> None:
        return None


class FakeRepository:
    """Simula tfm.documents y tfm.chunks con la sintaxis de filtros de PostgREST."""

    def __init__(self) -> None:
        self.documents: dict[str, dict[str, Any]] = {}
        self.chunks: list[dict[str, Any]] = []
        self.status_history: dict[str, list[str]] = {}
        self.fail_insert = False
        self.healthy = True

    def add_document(self, **values: Any) -> dict[str, Any]:
        doc = {
            "id": str(uuid4()),
            "title": "Documento de prueba",
            "area": "operaciones",
            "source_type": "procedimiento",
            "status": "pending",
            "chunk_count": 0,
            "error_message": None,
            "updated_at": None,
            "metadata": {},
            **values,
        }
        doc.setdefault("storage_path", f"{doc['area']}/{doc['id']}/doc.md")
        self.documents[doc["id"]] = doc
        return doc

    @staticmethod
    def _eq(params: dict[str, str], key: str) -> str | None:
        value = params.get(key)
        return value[3:] if value and value.startswith("eq.") else None

    async def select(self, table: str, params: dict[str, str] | None = None, *, user_jwt: str | None = None):
        params = params or {}
        if table != "documents":
            return []
        rows = list(self.documents.values())
        if doc_id := self._eq(params, "id"):
            rows = [r for r in rows if r["id"] == doc_id]
        if area := self._eq(params, "area"):
            rows = [r for r in rows if r["area"] == area]
        if sha := self._eq(params, "metadata->>sha256"):
            rows = [r for r in rows if (r.get("metadata") or {}).get("sha256") == sha]
        return [dict(r) for r in rows]

    async def insert(self, table: str, rows: Any, *, upsert_on: str | None = None, returning: bool = True):
        if self.fail_insert:
            raise RuntimeError("PostgREST caído")
        rows = rows if isinstance(rows, list) else [rows]
        if table == "documents":
            for row in rows:
                self.add_document(**row)
        elif table == "chunks":
            self.chunks.extend(rows)
        return rows if returning else []

    async def update(self, table: str, filters: dict[str, str], values: dict[str, Any]):
        doc = self.documents[self._eq(filters, "id")]
        doc.update(values)
        if "status" in values:
            self.status_history.setdefault(doc["id"], []).append(values["status"])
        return [dict(doc)]

    async def delete(self, table: str, filters: dict[str, str]) -> None:
        doc_id = self._eq(filters, "document_id")
        self.chunks = [c for c in self.chunks if c["document_id"] != doc_id]

    async def health(self) -> bool:
        return self.healthy

    async def aclose(self) -> None:
        return None


@pytest.fixture
def repo() -> FakeRepository:
    return FakeRepository()


@pytest.fixture
def storage() -> FakeStorage:
    return FakeStorage()


@pytest.fixture
def embedder() -> FakeEmbedder:
    return FakeEmbedder()


@pytest.fixture
def pipeline(repo, storage, embedder) -> IngestionPipeline:
    return IngestionPipeline(repo, storage, RecursiveChunker(300, 50), BatchEmbedder(embedder, batch_size=2), insert_batch_size=2)


@pytest.fixture
async def redis_client():
    client = fakeredis.FakeAsyncRedis(decode_responses=True)
    yield client
    await client.aclose()


@pytest.fixture
def queue(redis_client) -> IngestQueue:
    return IngestQueue(redis_client, "documents.ingest", "ingest-workers")


def store_document(repo: FakeRepository, storage: FakeStorage, content: bytes, filename: str, **values: Any) -> dict[str, Any]:
    """Registra un documento pendiente con su archivo en el Storage simulado."""
    doc = repo.add_document(metadata={"original_filename": filename}, **values)
    doc["storage_path"] = f"{doc['area']}/{doc['id']}/{filename}"
    storage.objects[f"documents/{doc['storage_path']}"] = content
    return doc
